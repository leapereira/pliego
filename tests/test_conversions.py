"""Exercise real artifacts and security boundaries, with external-engine skips."""
from pathlib import Path
import socket
import tempfile
import unittest
from unittest.mock import patch

import pymupdf as fitz

from server import conversions as conversion


class ConversionTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.outdir = self.root / "outputs"
        self.outdir.mkdir()
        self.pdf = self.root / "example.pdf"
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((60, 70), "Annual report 2026", fontsize=18)
            page.insert_text((60, 100), "Revenue increased to 120 units.")
            page.insert_text((60, 125), "=HYPERLINK(unsafe)")
            document.save(self.pdf)

    def tearDown(self):
        self.temporary.cleanup()

    def convert(self, tool, options=None, inputs=None, settings=None):
        return conversion.process_conversion(tool, [self.pdf] if inputs is None else inputs, options or {}, self.outdir, settings or {})

    def test_pdf_to_word_produces_editable_text(self):
        if not conversion.detect_engines()["pdf2docx"]:
            self.skipTest("pdf2docx is not installed")
        from docx import Document
        paths = self.convert("pdf-to-word")
        doc = Document(paths[0])
        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        self.assertIn("Annual report 2026", text)
        self.assertIn("120", text)

    def test_pdf_to_powerpoint_has_editable_text_and_images(self):
        from PIL import Image
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        image = self.root / "image.png"
        Image.new("RGB", (50, 50), "navy").save(image)
        pdf = self.root / "with-image.pdf"
        with fitz.open(self.pdf) as document:
            document[0].insert_image(fitz.Rect(60, 200, 160, 300), filename=image)
            document.save(pdf)
        result = self.convert("pdf-to-powerpoint", inputs=[pdf])[0]
        presentation = Presentation(result)
        self.assertEqual(len(presentation.slides), 1)
        shapes = presentation.slides[0].shapes
        text = "\n".join(shape.text for shape in shapes if shape.has_text_frame)
        self.assertIn("Annual report 2026", text)
        self.assertTrue(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in shapes))

    def test_scanned_powerpoint_explains_image_fallback(self):
        from PIL import Image, ImageDraw
        from pptx import Presentation
        image = Image.new("RGB", (500, 500), "white")
        ImageDraw.Draw(image).text((25, 25), "Scanned page", fill="black")
        image_path = self.root / "scan.png"
        image.save(image_path)
        pdf = self.root / "scan.pdf"
        with fitz.open() as document:
            page = document.new_page(width=500, height=500)
            page.insert_image(page.rect, filename=image_path)
            document.save(pdf)
        result = self.convert("pdf-to-powerpoint", inputs=[pdf])[0]
        presentation = Presentation(result)
        self.assertEqual(len(presentation.slides[0].shapes), 1)
        self.assertIn("sin texto editable", presentation.slides[0].notes_slide.notes_text_frame.text)

    def test_excel_fallback_and_formula_injection_prevention(self):
        from openpyxl import load_workbook
        result = self.convert("pdf-to-excel")[0]
        workbook = load_workbook(result)
        cells = list(workbook.active.values)
        self.assertTrue(any("120" in str(cell) for row in cells for cell in row))
        formula_cells = [cell for row in workbook.active for cell in row if "HYPERLINK" in str(cell.value)]
        self.assertTrue(formula_cells)
        self.assertTrue(all(cell.data_type != "f" for cell in formula_cells))
        self.assertTrue(all(cell.value.startswith("'=") for cell in formula_cells))

    def test_excel_extracts_table_columns(self):
        from openpyxl import load_workbook
        table_pdf = self.root / "table.pdf"
        with fitz.open() as document:
            page = document.new_page()
            for x in (50, 200, 350):
                page.draw_line((x, 60), (x, 180))
            for y in (60, 100, 140, 180):
                page.draw_line((50, y), (350, y))
            for x, y, value in ((60, 85, "Product"), (210, 85, "Units"), (60, 125, "Alpha"), (210, 125, "12"), (60, 165, "Beta"), (210, 165, "24")):
                page.insert_text((x, y), value)
            document.save(table_pdf)
        result = self.convert("pdf-to-excel", inputs=[table_pdf])[0]
        workbook = load_workbook(result)
        self.assertEqual(workbook.active["A1"].value, "Product")
        self.assertEqual(workbook.active["B2"].value, "12")

    def test_html_renders_text_and_table_without_active_resources(self):
        raw = '<h1>Hola mundo</h1><p>España y Uruguay</p><script>PRIVATE_SCRIPT</script><iframe src="file:///etc/passwd">PRIVATE_FRAME</iframe><img src="file:///etc/passwd"><table><tr><td>Valor</td><td>123</td></tr></table>'
        result = self.convert("html-to-pdf", {"html": raw}, inputs=[])[0]
        with fitz.open(result) as document:
            text = "".join(page.get_text() for page in document)
        self.assertIn("Hola mundo", text)
        self.assertIn("123", text)
        self.assertNotIn("PRIVATE_SCRIPT", text)
        self.assertNotIn("PRIVATE_FRAME", text)
        self.assertNotIn("root:", text)

    def test_html_strips_all_resource_attributes(self):
        safe = conversion._safe_html('<div style="background:url(file:///etc/passwd)" onclick="fetch(1)"><a href="https://example.com">Texto</a><object data="file:///x">oculto</object></div>')
        self.assertNotIn("file:", safe)
        self.assertNotIn("style=", safe)
        self.assertNotIn("onclick", safe)
        self.assertNotIn("href=", safe)
        self.assertIn("Texto", safe)

    def test_private_metadata_and_non_http_urls_rejected(self):
        for url in ("file:///etc/passwd", "http://127.0.0.1", "http://169.254.169.254/latest/meta-data", "http://10.0.0.1", "http://[::1]", "https://example.com:8443", "https://user:password@example.com"):
            with self.subTest(url=url), self.assertRaises(ValueError):
                conversion._public_url(url)

    def test_dns_mixed_public_and_private_addresses_rejected(self):
        records = [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443)), (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("127.0.0.1", 443))]
        with patch("server.conversions.socket.getaddrinfo", return_value=records):
            with self.assertRaises(ValueError):
                conversion._public_url("https://example.com")

    def test_url_validation_pins_public_dns_resolution(self):
        records = [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))]
        with patch("server.conversions.socket.getaddrinfo", return_value=records):
            parts = conversion._public_url("https://example.com/path?x=1")
        self.assertEqual(parts, ("https", "example.com", 443, "/path?x=1", "93.184.216.34"))

    def test_ai_requires_explicit_provider_and_model(self):
        with self.assertRaisesRegex(ValueError, "Configura un proveedor"):
            self.convert("summarize")

    def test_summary_uses_all_fragments_without_silent_truncation(self):
        fragments = []

        def response(settings, instruction, content):
            fragments.append(content)
            return "Resumen con 120 unidades."

        text = " ".join(f"Token{i}" for i in range(6000))
        with patch("server.conversions._pdf_text", return_value=text), patch("server.conversions._ai_request", side_effect=response):
            result = self.convert("summarize", settings={"ai_base_url": "http://localhost:11434/v1", "ai_model": "local-model"})[0]
        self.assertIn("Resumen", result.read_text())
        self.assertEqual("".join(fragments[:-1]), text)

    def test_translation_writes_readable_pdf_and_keeps_all_text(self):
        with patch("server.conversions._ai_request", return_value="Informe anual 2026\nIngresos de 120 unidades."):
            result = self.convert("translate", {"target_language": "español"}, settings={"ai_base_url": "http://localhost:11434/v1", "ai_model": "local-model"})[0]
        with fitz.open(result) as document:
            self.assertIn("Ingresos de 120 unidades", "".join(page.get_text() for page in document))
            self.assertIn("refluido", document.metadata["title"])

    def test_ai_reports_output_truncation(self):
        import httpx
        request = httpx.Request("POST", "http://localhost:11434/v1/chat/completions")
        response = httpx.Response(200, request=request, json={"choices": [{"finish_reason": "length", "message": {"content": "incomplete"}}]})
        with patch("server.conversions.httpx.post", return_value=response), self.assertRaisesRegex(ValueError, "cortó"):
            conversion._ai_request({"ai_base_url": "http://localhost:11434/v1", "ai_model": "example"}, "Translate", "Text")

    def test_real_office_word_excel_and_powerpoint_to_pdf(self):
        if not conversion.detect_engines()["libreoffice"]:
            self.skipTest("LibreOffice is not installed")
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches

        word = self.root / "letter.docx"
        doc = Document()
        doc.add_paragraph("Office document conversion works")
        doc.save(word)
        excel = self.root / "sheet.xlsx"
        workbook = Workbook()
        workbook.active.append(["Product", "Units"])
        workbook.active.append(["Alpha", 12])
        workbook.save(excel)
        powerpoint = self.root / "slides.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Presentation conversion works"
        presentation.save(powerpoint)
        for tool, source, expected in (("word-to-pdf", word, "Office document"), ("excel-to-pdf", excel, "Product"), ("powerpoint-to-pdf", powerpoint, "Presentation conversion")):
            with self.subTest(tool=tool):
                result = self.convert(tool, inputs=[source])[0]
                with fitz.open(result) as document:
                    text = "".join(page.get_text() for page in document)
                self.assertIn(expected, text)

    def test_real_tesseract_creates_searchable_pdf(self):
        if not conversion.detect_engines()["tesseract"]:
            self.skipTest("Tesseract is not installed")
        spanish = self.root / "espanol.pdf"
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((60, 70), "Documento en espanol", fontsize=22)
            page.insert_text((60, 110), "El informe contiene 120 unidades.", fontsize=16)
            document.save(spanish)
        for language, source, expected in (("eng", self.pdf, "Annual"), ("spa", spanish, "Documento")):
            with self.subTest(language=language):
                result = self.convert("ocr", {"language": language, "dpi": 200}, inputs=[source])[0]
                with fitz.open(result) as document:
                    self.assertEqual(document.page_count, 1)
                    self.assertIn(expected, document[0].get_text())

    def test_real_ghostscript_declares_pdfa2_output_intent(self):
        executable = conversion._engine("ghostscript")
        if not executable or not conversion._srgb_profile(executable):
            self.skipTest("Ghostscript and an sRGB ICC profile are required")
        result = self.convert("pdf-to-pdfa")[0]
        with fitz.open(result) as document:
            self.assertIn("pdfaid:part", document.get_xml_metadata())
            self.assertIn("Annual report", document[0].get_text())
            catalog = document.xref_object(document.pdf_catalog())
            self.assertIn("OutputIntents", catalog)


if __name__ == "__main__":
    unittest.main()
